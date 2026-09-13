"""Reference folders: local files JARVIS reads in place as confidential reference.

Files are never copied, cached on disk or modified. Listing uses file metadata
only; text is extracted in memory when a search needs it and cached per
(path, size, mtime) for the life of the process.

Classification comes from the folder's configured level. A top-level subfolder
named after a higher level (e.g. `SECRET/`) raises it; nothing can lower it.
Hidden files, Office lock files and symlinks are ignored, so a link cannot pull
files from outside the folder into JARVIS.
"""
from __future__ import annotations

import hashlib
import os
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from .config import ReferenceFolder
from .private_brain import Classification, KnowledgeDocument

ORIGIN = "reference_folder"
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".json"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".docx", ".pptx", ".pdf"}
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_FILES_PER_FOLDER = 1000
MAX_CHARS_PER_FILE = 500_000


class ExtractionError(Exception):
    """A file could not be turned into text (encrypted, corrupt, missing library)."""


def _document_id(path: Path) -> str:
    return "ref-" + hashlib.sha256(str(path).encode("utf-8")).hexdigest()[:16]


def _classification_for(root: ReferenceFolder, relative: Path) -> Classification:
    level = root.classification
    if len(relative.parts) > 1:
        try:
            level = max(level, Classification.parse(relative.parts[0]))
        except ValueError:
            pass
    return level


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        raise ExtractionError("python-docx is not installed") from None
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        raise ExtractionError("python-pptx is not installed") from None
    parts = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        if texts:
            parts.append(f"Slide {number}\n" + "\n".join(texts))
    return "\n\n".join(parts)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ExtractionError("pypdf is not installed") from None
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ExtractionError("PDF is password-protected")
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8", errors="replace")
        elif suffix == ".docx":
            text = _extract_docx(path)
        elif suffix == ".pptx":
            text = _extract_pptx(path)
        elif suffix == ".pdf":
            text = _extract_pdf(path)
        else:
            raise ExtractionError(f"unsupported file type {suffix}")
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"could not read file ({type(e).__name__})") from None
    return text[:MAX_CHARS_PER_FILE]


class ReferenceLibrary:
    def __init__(self, folders: Iterable[ReferenceFolder], owner: str):
        self.folders = tuple(folders)
        self.owner = owner
        self._paths: dict[str, Path] = {}
        self._cache: dict[str, tuple[tuple[int, int], str]] = {}
        self._lock = threading.Lock()

    def _scan(self, folder: ReferenceFolder):
        root = folder.path
        if not root.is_dir():
            return
        count = 0
        for dirpath, dirnames, filenames in os.walk(root, followlinks=False):
            dirnames[:] = sorted(d for d in dirnames
                                 if not d.startswith(".") and not os.path.islink(os.path.join(dirpath, d)))
            for name in sorted(filenames):
                if name.startswith((".", "~$")):
                    continue
                path = Path(dirpath) / name
                if path.suffix.lower() not in SUPPORTED_SUFFIXES or path.is_symlink():
                    continue
                try:
                    st = path.stat()
                except OSError:
                    continue
                if st.st_size > MAX_FILE_BYTES:
                    continue
                yield path, st
                count += 1
                if count >= MAX_FILES_PER_FOLDER:
                    return

    def documents(self) -> list[KnowledgeDocument]:
        docs, paths = [], {}
        for folder in self.folders:
            for path, st in self._scan(folder):
                relative = path.relative_to(folder.path)
                doc_id = _document_id(path)
                paths[doc_id] = path
                docs.append(KnowledgeDocument(
                    document_id=doc_id,
                    title=path.stem,
                    classification=_classification_for(folder, relative),
                    owner=self.owner,
                    allowed_roles=folder.allowed_roles,
                    metadata={
                        "origin": ORIGIN,
                        "source": str(relative),
                        "version": "1",
                        "ingested_at": datetime.fromtimestamp(st.st_mtime, timezone.utc)
                        .isoformat(timespec="seconds"),
                    },
                ))
        with self._lock:
            self._paths = paths
            for stale in set(self._cache) - set(paths):
                del self._cache[stale]
        return docs

    def text(self, document: KnowledgeDocument) -> str:
        with self._lock:
            path = self._paths.get(document.document_id)
        if path is None:
            raise ExtractionError("reference file is no longer available")
        try:
            st = path.stat()
        except OSError:
            raise ExtractionError("reference file is no longer available") from None
        key = (st.st_size, st.st_mtime_ns)
        with self._lock:
            cached = self._cache.get(document.document_id)
            if cached and cached[0] == key:
                return cached[1]
        text = extract_text(path)
        with self._lock:
            self._cache[document.document_id] = (key, text)
        return text


def is_reference(document: KnowledgeDocument) -> bool:
    return document.metadata.get("origin") == ORIGIN


__all__ = ["ExtractionError", "ReferenceLibrary", "extract_text", "is_reference"]
