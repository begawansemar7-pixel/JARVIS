"""PowerPoint (.pptx) board deck writer — 16:9, drawn on the blank layout."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import theme
from .model import ExecutiveDocument, Slide, paginate_slides

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.7)


def _rgb(color) -> RGBColor:
    return RGBColor(*color)


def _rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    return shape


def _text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = theme.FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return box


def _footer(slide, document: ExecutiveDocument, number: int) -> None:
    y = H - Inches(0.5)
    _rect(slide, MARGIN, y - Inches(0.08), W - 2 * MARGIN, Emu(9525), theme.RULE)
    _text(slide, MARGIN, y, Inches(6), Inches(0.3), document.title, 10, theme.MUTED)
    _text(slide, Inches(5.2), y, Inches(3), Inches(0.3), theme.marking(document.classification),
          10, theme.GOLD, bold=True, align=PP_ALIGN.CENTER)
    _text(slide, W - MARGIN - Inches(1), y, Inches(1), Inches(0.3), str(number), 10, theme.MUTED,
          align=PP_ALIGN.RIGHT)


def _title_slide(prs, document: ExecutiveDocument) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, W, H, theme.NAVY)
    _text(slide, MARGIN, Inches(1.2), W - 2 * MARGIN, Inches(2.0), document.title, 40, theme.WHITE,
          bold=True, anchor=MSO_ANCHOR.BOTTOM)
    _rect(slide, MARGIN, Inches(3.55), Inches(1.2), Inches(0.06), theme.GOLD)
    if document.subtitle:
        _text(slide, MARGIN, Inches(3.85), W - 2 * MARGIN, Inches(0.8), document.subtitle, 20, theme.GOLD)
    meta = " · ".join(x for x in (document.author, document.date) if x)
    _text(slide, MARGIN, H - Inches(1.3), Inches(8), Inches(0.4), meta, 14, theme.WHITE)
    _text(slide, W - MARGIN - Inches(4), H - Inches(1.3), Inches(4), Inches(0.4),
          theme.marking(document.classification), 14, theme.GOLD, bold=True, align=PP_ALIGN.RIGHT)


def _bullets(slide, x, y, w, h, bullets: list[str]) -> None:
    frame = slide.shapes.add_textbox(x, y, w, h).text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    size = 20 if len(bullets) <= 4 else 17
    for i, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(10)
        marker = p.add_run()
        marker.text = "■  "
        marker.font.size = Pt(size - 6)
        marker.font.color.rgb = _rgb(theme.GOLD)
        run = p.add_run()
        run.text = bullet
        run.font.name = theme.FONT
        run.font.size = Pt(size)
        run.font.color.rgb = _rgb(theme.INK)


def _table(slide, x, y, w, table) -> None:
    rows, cols = len(table.rows) + 1, len(table.columns)
    shape = slide.shapes.add_table(rows, cols, x, y, w, Inches(0.45) * rows)
    grid = shape.table
    for c, name in enumerate(table.columns):
        cell = grid.cell(0, c)
        cell.text = name
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(theme.NAVY)
        font = cell.text_frame.paragraphs[0].runs[0].font
        font.name, font.size, font.bold = theme.FONT, Pt(13), True
        font.color.rgb = _rgb(theme.WHITE)
    for r, row in enumerate(table.rows, start=1):
        for c, value in enumerate(row):
            cell = grid.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(theme.TABLE_STRIPE if r % 2 == 0 else theme.WHITE)
            font = cell.text_frame.paragraphs[0].runs[0].font
            font.name, font.size = theme.FONT, Pt(12)
            font.color.rgb = _rgb(theme.INK)


def _content_slide(prs, document: ExecutiveDocument, data: Slide, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, Inches(0.18), H, theme.NAVY)
    _text(slide, MARGIN, Inches(0.45), W - 2 * MARGIN, Inches(0.4), data.title.upper(), 13, theme.GOLD, bold=True)
    top = Inches(0.9)
    if data.message:
        _text(slide, MARGIN, top, W - 2 * MARGIN, Inches(1.3), data.message, 28, theme.NAVY, bold=True)
        top = Inches(2.3)
    body_h = H - top - Inches(0.9)
    if data.table:
        _table(slide, MARGIN, top, W - 2 * MARGIN, data.table)
    elif data.bullets:
        _bullets(slide, MARGIN, top, W - 2 * MARGIN, body_h, data.bullets)
    if data.notes:
        slide.notes_slide.notes_text_frame.text = data.notes
    _footer(slide, document, number)


def write_pptx(document: ExecutiveDocument, path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    _title_slide(prs, document)
    for number, data in enumerate(paginate_slides(document.slides), start=2):
        _content_slide(prs, document, data, number)
    prs.core_properties.title = document.title
    prs.core_properties.author = document.author or "JARVIS"
    prs.core_properties.category = theme.marking(document.classification)
    prs.save(path)
    return path
