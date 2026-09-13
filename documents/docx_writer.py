"""Word (.docx) report writer."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import theme
from .model import ExecutiveDocument, Table


def _rgb(color) -> RGBColor:
    return RGBColor(*color)


def _shade(cell, color) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), theme.hex_color(color)[1:])
    cell._tc.get_or_add_tcPr().append(shd)


def _style_fonts(doc: Document) -> None:
    for name, size, color, bold in (
        ("Normal", 11, theme.INK, False),
        ("Title", 28, theme.NAVY, True),
        ("Subtitle", 14, theme.MUTED, False),
        ("Heading 1", 16, theme.NAVY, True),
        ("Heading 2", 13, theme.NAVY, True),
        ("List Bullet", 11, theme.INK, False),
    ):
        style = doc.styles[name]
        style.font.name = theme.FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = _rgb(color)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), theme.FONT)
    doc.styles["List Bullet"].paragraph_format.space_after = Pt(3)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)


def _table(doc: Document, table: Table) -> None:
    t = doc.add_table(rows=1, cols=len(table.columns))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = doc.styles["Table Grid"]
    for i, name in enumerate(table.columns):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(name)
        run.bold = True
        run.font.color.rgb = _rgb(theme.WHITE)
        _shade(cell, theme.NAVY)
    for r, row in enumerate(table.rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            if r % 2:
                _shade(cells[i], theme.TABLE_STRIPE)
    doc.add_paragraph()


def write_docx(document: ExecutiveDocument, path: Path) -> Path:
    doc = Document()
    _style_fonts(doc)
    marking = theme.marking(document.classification)

    section = doc.sections[0]
    for part, align in ((section.header, WD_ALIGN_PARAGRAPH.RIGHT), (section.footer, WD_ALIGN_PARAGRAPH.CENTER)):
        p = part.paragraphs[0]
        p.alignment = align
        run = p.add_run(marking)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = _rgb(theme.GOLD)

    doc.add_paragraph(document.title, style="Title")
    if document.subtitle:
        doc.add_paragraph(document.subtitle, style="Subtitle")
    meta = " · ".join(x for x in (document.author, document.date, marking) if x)
    doc.add_paragraph(meta).runs[0].font.color.rgb = _rgb(theme.MUTED)

    if document.summary:
        doc.add_heading("Executive Summary", level=1)
        for line in document.summary:
            doc.add_paragraph(line, style="List Bullet")

    for sec in document.sections:
        doc.add_heading(sec.heading, level=1)
        for para in sec.paragraphs:
            doc.add_paragraph(para)
        for bullet in sec.bullets:
            doc.add_paragraph(bullet, style="List Bullet")
        if sec.table:
            _table(doc, sec.table)

    doc.core_properties.title = document.title
    doc.core_properties.author = document.author or "JARVIS"
    doc.core_properties.category = marking
    doc.save(path)
    return path
