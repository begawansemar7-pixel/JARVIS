import json
from datetime import datetime

import pytest
from docx import Document
from pptx import Presentation

from actions import ceo_documents as action
from core import undo
from core.action_loader import _validate
from documents import decision_to_document, generate
from documents.generator import ENV_OUTPUT_DIR, output_dir, parse_formats, slugify
from documents.model import DATA_GAP, ExecutiveDocument, Section, Slide, paginate_slides, parse_slides

DECISION = {
    "decision_question": "Should we pilot an AI data center in Kalimantan?",
    "decision_owner": "CEO",
    "options": [
        {"name": "A", "description": "Build 150MW", "scores": {"fit": 9, "risk": 4}},
        {"name": "B", "description": "Pilot 30MW", "scores": {"fit": 8, "risk": 7}},
    ],
    "strategic_phase": {"pattern": "升 Sheng — gradual scaling", "rationale": "Scale with demand"},
    "recommendation": {"decision": "Pilot 30MW with a month-9 gate", "why": "Keeps optionality"},
    "timing": {"classification": "PILOT_NOW"},
}


@pytest.fixture
def out_dir(tmp_path, monkeypatch):
    monkeypatch.setenv(ENV_OUTPUT_DIR, str(tmp_path))
    return tmp_path


def sample_document():
    return ExecutiveDocument(
        title="Quarterly Board Report",
        classification="CONFIDENTIAL",
        summary=["Revenue up 12%"],
        sections=[Section("Performance", paragraphs=["Solid quarter."], bullets=["EBITDA 48%"])],
        slides=[Slide("Headline", "Growth is accelerating", ["Revenue up 12%"])],
    )


def test_default_output_dir_is_documents_jarvis(monkeypatch):
    monkeypatch.delenv(ENV_OUTPUT_DIR, raising=False)
    assert output_dir().parts[-2:] == ("Documents", "JARVIS")


def test_format_parsing():
    assert parse_formats(None, "report") == ("docx", "pdf")
    assert parse_formats("doc, PDF", "report") == ("docx", "pdf")
    with pytest.raises(ValueError):
        parse_formats("xlsx", "report")
    with pytest.raises(ValueError):
        parse_formats("pptx", "report")
    with pytest.raises(ValueError):
        parse_formats("pdf", "memo")


def test_slugify_is_filesystem_safe():
    assert slugify("Ekspansi / Data: Center — AI?") == "Ekspansi-Data-Center-AI"
    assert slugify("升") == "Document"


def test_generates_all_formats_readable(tmp_path):
    files = generate(sample_document(), "both", directory=tmp_path, now=datetime(2026, 9, 13, 9, 30))
    assert {(f.kind, f.format) for f in files} == {
        ("report", "docx"), ("report", "pdf"), ("presentation", "pptx"), ("presentation", "pdf"),
    }
    by = {(f.kind, f.format): f.path for f in files}
    text = "\n".join(p.text for p in Document(by[("report", "docx")]).paragraphs)
    assert "Quarterly Board Report" in text and "EBITDA 48%" in text
    assert len(Presentation(by[("presentation", "pptx")]).slides) == 2  # title + 1
    for key in (("report", "pdf"), ("presentation", "pdf")):
        assert by[key].read_bytes().startswith(b"%PDF")
    assert by[("report", "docx")].name == "2026-09-13_0930_CEO-Report_Quarterly-Board-Report.docx"


def test_never_overwrites_existing_files(tmp_path):
    now = datetime(2026, 9, 13, 9, 30)
    first = generate(sample_document(), "report", "docx", directory=tmp_path, now=now)[0].path
    second = generate(sample_document(), "report", "docx", directory=tmp_path, now=now)[0].path
    assert first != second and first.exists() and second.exists()


def test_empty_presentation_is_rejected(tmp_path):
    with pytest.raises(ValueError):
        generate(ExecutiveDocument(title="Empty"), "presentation", directory=tmp_path)
    assert list(tmp_path.iterdir()) == []


def test_overfull_slides_are_split():
    slides = paginate_slides([Slide("Risks", "msg", [f"risk {i}" for i in range(14)])])
    assert [len(s.bullets) for s in slides] == [6, 6, 2]
    assert slides[1].title.endswith("(cont.)") and slides[1].message == ""


def test_decision_mapping_marks_gaps_instead_of_inventing():
    doc = decision_to_document(DECISION, classification="confidential")
    assert doc.classification == "CONFIDENTIAL"
    assert doc.title == DECISION["decision_question"]
    assert any(DATA_GAP in b for s in doc.sections for b in s.bullets)
    comparison = next(s for s in doc.slides if s.title == "Option comparison")
    assert comparison.table.rows == [["A", "9", "4"], ["B", "8", "7"]]
    assert 12 <= len(doc.slides) <= 18


def test_cjk_content_renders_to_pdf(tmp_path):
    doc = decision_to_document(DECISION)
    files = generate(doc, "presentation", "pdf", directory=tmp_path)
    assert files[0].path.stat().st_size > 1000


def test_slides_accept_json_strings():
    slides = parse_slides(json.dumps([{"title": "One", "bullets": "- a\n- b"}, {"bullets": ["no title"]}]))
    assert len(slides) == 1 and slides[0].bullets == ["a", "b"]


def test_action_contract_is_valid():
    record = _validate(action, "ceo_documents.py")
    assert record.valid, record.error


def test_action_creates_files_and_undo_removes_them(out_dir):
    undo.clear()
    out = action.ceo_document({
        "document_type": "presentation",
        "decision_json": json.dumps(DECISION),
        "classification": "CONFIDENTIAL",
    })
    assert out.startswith("Saved 2 file(s)")
    created = sorted(p.suffix for p in out_dir.iterdir())
    assert created == [".pdf", ".pptx"]
    assert undo.can_undo()
    undo.undo_last()
    assert list(out_dir.iterdir()) == []


def test_action_reports_invalid_input_without_writing(out_dir):
    assert action.ceo_document({"document_type": "report"}).startswith("Document not created")
    assert action.ceo_document({"document_type": "report", "title": "X", "classification": "ULTRA"}) \
        .startswith("Document not created")
    assert action.ceo_document({"document_type": "report", "decision_json": "{oops"}) \
        .startswith("Document not created")
    assert list(out_dir.iterdir()) == []
